import pytest

from eri.cleanup import clean_text
from eri.passages import load_passages
from eri.spellcheck import correct_text
from tests.data.patent_examples import FROG, MARKET


def fix(kiwi, text):
    log = []
    return correct_text(text, kiwi, log), log


@pytest.mark.parametrize("wrong, right", [
    ("그래서 늘 가난 하고 고독했다.", "그래서 늘 가난하고 고독했다."),
    ("그래서 늘 가 난하고 고독했다.", "그래서 늘 가난하고 고독했다."),
    ("학교 에서 친구들 과 놀았다.", "학교에서 친구들과 놀았다."),
    ("이 그 림은 걸작으로 손꼽힌다.", "이 그림은 걸작으로 손꼽힌다."),
    ("자신만의 채 색 기법을 만들었다.", "자신만의 채색 기법을 만들었다."),
    ("그는 그것을알수있다고 했다.", "그는 그것을 알 수 있다고 했다."),
    ("그렇게 됬다. 몇일 동안 기다렸어요. 내일 할께요.", "그렇게 됐다. 며칠 동안 기다렸어요. 내일 할게요."),
])
def test_corrections(kiwi, wrong, right):
    text, log = fix(kiwi, wrong)
    assert text == right
    assert log and all(c.before != c.after for c in log)


@pytest.mark.parametrize("ok", [
    "작은 집에 살았다. 큰 일이 났다.", "이 그림은 걸작이다.", "할 수 있다.", "그 사람 이 책을 읽었다.",
    "국어교육을 먹고있다.", "드라마 〈오징어 게임〉은 인기가 많았다.", "크지 않되 단단하다.",
])
def test_correct_text_is_left_alone(kiwi, ok):
    assert fix(kiwi, ok) == (ok, [])


def test_no_false_corrections_on_clean_passages(kiwi):
    texts = [p.text for p in load_passages("지문모음.txt")] + [FROG, MARKET]
    changed = []
    for t in texts:
        _, log = fix(kiwi, t)
        changed += [(c.before, c.after) for c in log]
    # '김 매기' → '김매기'(사전상 한 단어)만 고친다
    assert changed == [("고기잡이, 김 매기 등", "고기잡이, 김매기 등")]


def test_cleanup_logs_line_joins_and_markers(kiwi):
    log = []
    clean_text("얼핏 보기에 아주 어\n둡고 칙칙해 보이는 ㉠ 그림이다. 그리고 다른\n문장도 있다.", kiwi, log=log)
    kinds = [c.kind for c in log]
    assert kinds.count("기호 삭제") == 1 and kinds.count("줄바꿈 복원") >= 1


def test_analysis_records_corrections(kiwi, mini_vocab):
    from eri.analyzer import ERIAnalyzer
    from eri.config import ERIConfig
    from eri.passages import Passage
    cfg = ERIConfig()
    r = ERIAnalyzer(mini_vocab, cfg, kiwi).analyze(Passage("시험", "개구리 들은 겨울 에 잠을 자. 물론 아무것도 먹지 않아.", "초등"))
    assert r.corrected_text.startswith("개구리들은 겨울에 잠을 자.")
    assert len(r.corrections) == 2
    cfg.auto_correct = False
    r = ERIAnalyzer(mini_vocab, cfg, kiwi).analyze(Passage("시험", "개구리 들은 겨울 에 잠을 자.", "초등"))
    assert r.corrections == [] and r.corrected_text == "개구리 들은 겨울 에 잠을 자."


def test_report_has_correction_sheets(tmp_path, kiwi, mini_vocab):
    from openpyxl import load_workbook

    from eri.analyzer import ERIAnalyzer
    from eri.config import ERIConfig
    from eri.passages import Passage
    from eri.report import write_report
    cfg = ERIConfig()
    r = ERIAnalyzer(mini_vocab, cfg, kiwi).analyze(Passage("시험", "개구리 들은 겨울 에 잠을 자.", "초등"))
    out = tmp_path / "r.xlsx"
    write_report([r], cfg, out, "mini")
    wb = load_workbook(out, rich_text=True)
    assert {"교정 내역", "교정된 지문"} <= set(wb.sheetnames)
    rows = list(wb["교정 내역"].iter_rows(min_row=2, max_row=3, values_only=True))
    assert [row[2] for row in rows] == ["띄어쓰기(붙임)", "띄어쓰기(붙임)"]
    assert str(rows[0][4]).startswith("개구리들은")
