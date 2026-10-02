from openpyxl import Workbook, load_workbook

from eri.config import ERIConfig
from eri.passages import load_passages, parse_text
from eri.qualitative import pairwise_agreement, parse_score
from eri.vocab import find_vocab_file, load_vocabulary


def _xlsx(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def test_vocab_autodetect_headers_sheets_and_homographs(tmp_path):
    _xlsx(tmp_path / "아무 이름.xlsx", {
        "설명": [["이 시트는 안내문입니다"]],
        "1등급": [["국어 기초 어휘"], ["등급 ", "어휘", "품사"], ["1등급", "가정", "명사"], ["1등급", "가다", "동사"]],
        "4등급": [["등급", "어휘", "품사"], ["4등급", "가정", "명사"], ["4등급", "가정하다", "동사"]],
    })
    _xlsx(tmp_path / "ERI_분석결과.xlsx", {"ERI 결과": [["지문명", "ERI"]]})
    found = find_vocab_file([tmp_path])
    assert found.name == "아무 이름.xlsx"
    v = load_vocabulary(found, use_cache=False)
    assert v.grade("가정") == 1                   # 마지막 행(4등급)으로 덮어쓰지 않음
    assert v.grade("가정하다") == 4


def test_vocab_csv_and_letter_grades(tmp_path):
    p = tmp_path / "words.csv"
    p.write_text("단어,등급\n사과,A\n사과나무,B등급\n과수원,C\n", encoding="cp949")
    v = load_vocabulary(p, use_cache=False)
    assert (v.grade("사과"), v.grade("사과나무"), v.grade("과수원")) == (1, 2, 3)


def test_legacy_colon_format_cp949_crlf(tmp_path):
    text = ("제주 민요: 제주도는 섬이다.\r\n예: 이 줄은 본문이다.\r\n\r\n"
            "[초등] 개구리의 겨울: 개구리는 겨울잠을 잔다.\r\n봄이 오면 깨어난다.\r\n")
    p = tmp_path / "지문.txt"
    p.write_bytes(text.encode("cp949"))
    ps = load_passages(p)
    assert [x.title for x in ps] == ["제주 민요", "개구리의 겨울"]
    assert "예: 이 줄은 본문이다." in ps[0].text
    assert ps[1].level == "초등" and ps[0].level is None


def test_original_passage_file_loads():
    ps = load_passages("지문모음.txt")
    assert len(ps) == 10
    assert ps[0].title == "판소리 <흥보가>의 장단과 표현 기법"


def test_heading_and_block_formats():
    ps = parse_text("# 첫 지문 [5학년]\n본문 하나.\n\n# 둘째 지문\n본문 둘.\n")
    assert [(p.title, p.level, p.grade) for p in ps] == [("첫 지문", "초등", 5.0), ("둘째 지문", None, None)]
    ps = parse_text("제목 없음\n첫 문장이다.\n\n\n두 번째 덩어리이다.\n이어진다.", default_title="파일")
    assert [p.title for p in ps] == ["제목 없음", "파일 2"]
    ps = parse_text("그냥 한 덩어리 본문이다. 제목이 없다.", default_title="내 파일")
    assert ps[0].title == "내 파일"


def test_table_and_folder_inputs(tmp_path):
    _xlsx(tmp_path / "지문표.xlsx", {"Sheet": [["번호", "제목", "본문", "학교급"],
                                              [1, "가", "첫 본문이다.", "초등"], [2, "나", "둘째 본문이다.", "중학교"]]})
    (tmp_path / "추가.txt").write_text("다른 지문: 내용이다.", encoding="utf-8")
    ps = load_passages(tmp_path)
    assert {(p.title, p.level) for p in ps} == {("가", "초등"), ("나", "중등"), ("다른 지문", None)}


def test_qualitative_scores_and_agreement():
    assert parse_score("") is None and parse_score(" +1 ") == 1.0
    try:
        parse_score("4")
        raise AssertionError
    except ValueError:
        pass
    mean, pairs = pairwise_agreement([[1, 0, -1], [1, 0, 0], [1, 0, -1]])
    assert pairs[(1, 3)] == 1.0 and round(mean, 3) == round((2 / 3 + 1 + 2 / 3) / 3, 3)


def test_report_and_recalc_roundtrip(tmp_path, kiwi, mini_vocab):
    from eri.analyzer import ERIAnalyzer
    from eri.passages import Passage
    from eri.report import recalc_report, write_report
    cfg = ERIConfig()
    an = ERIAnalyzer(mini_vocab, cfg, kiwi)
    r = an.analyze(Passage("시장", "시장은 사람이 모여 거래하는 곳이다. 가격이 오르면 수요량은 감소하게 된다.", "중등"))
    out = tmp_path / "결과.xlsx"
    write_report([r], cfg, out, "mini")
    wb = load_workbook(out)
    ws = wb["ERI 결과"]
    header = [c.value for c in ws[1]]
    for i, v in enumerate([1, 0, 0.5]):
        ws.cell(2, header.index(f"평가자{i + 1}") + 1).value = v
    wb.save(out)
    recalc_report(out, cfg)
    ws = load_workbook(out)["ERI 결과"]
    row = {h: c.value for h, c in zip(header, ws[2])}
    assert row["정성 지수(평균)"] == 0.5
    assert row["ERI"] == round(row["정량 지수"] + 0.5, 1)
    assert row["학년·단계"]


def _docx(path, body_xml):
    import zipfile
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", f"<w:document {ns}><w:body>{body_xml}</w:body></w:document>")


def _p(text):
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def _row(*cells):
    return "<w:tr>" + "".join(f"<w:tc>{''.join(_p(t) for t in c.split('|'))}</w:tc>" for c in cells) + "</w:tr>"


def test_docx_template_file_in_repo():
    ps = load_passages("견본/지문_입력_견본.docx")
    assert [(p.title, p.level) for p in ps] == [("개구리의 겨울은 잠자는 시간", "초등"), ("수요와 공급", "중등")]
    assert all(len(p.text.split()) > 90 for p in ps)


def test_docx_tables_merged_or_blank(tmp_path):
    # 표 두 개가 하나로 붙은 경우 + 내용이 빈 표 + 표 밖 안내문
    body = (_p("안내문: 이 글은 무시된다.")
            + "<w:tbl>" + _row("제목", "첫 지문") + _row("학교급|초등 / 중등", "초등") + _row("내용", "첫 문단.|둘째 문단.")
            + _row("제목:", "둘째 지문") + _row("학교급", "") + _row("내용", "본문이다.") + "</w:tbl>"
            + "<w:tbl>" + _row("제목", "") + _row("학교급", "") + _row("내용", "") + "</w:tbl>")
    _docx(tmp_path / "a.docx", body)
    ps = load_passages(tmp_path / "a.docx")
    assert [(p.title, p.level, p.text) for p in ps] == [
        ("첫 지문", "초등", "첫 문단.\n둘째 문단."), ("둘째 지문", None, "본문이다.")]


def test_docx_paragraph_labels(tmp_path):
    _docx(tmp_path / "b.docx", _p("제목: 가 지문") + _p("학교급: 중등") + _p("내용: 첫 줄이다.") + _p("이어진다.")
          + _p("제목: 나 지문") + _p("내용: 둘째다."))
    ps = load_passages(tmp_path / "b.docx")
    assert [(p.title, p.level) for p in ps] == [("가 지문", "중등"), ("나 지문", None)]
    assert ps[0].text == "첫 줄이다.\n이어진다."
