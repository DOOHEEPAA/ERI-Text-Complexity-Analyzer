"""어휘 등급 사전 불러오기.

- 파일 이름·시트 이름이 정확히 일치할 필요가 없다. '어휘'와 '등급' 열이 있는
  엑셀(.xlsx)/CSV 파일과 시트를 자동으로 찾는다.
- 같은 형태의 단어(동형어)가 여러 등급에 있으면 마지막 행으로 덮어쓰지 않고
  품사가 맞는 항목 중 가장 쉬운(낮은) 등급을 쓴다.
"""
from __future__ import annotations

import csv
import pickle
import re
from collections import defaultdict
from pathlib import Path

from .textutil import normalize_header, read_text

WORD_HEADERS = ("어휘", "단어", "표제어", "word")
GRADE_HEADERS = ("등급", "어휘등급", "수준", "grade", "level")
POS_HEADERS = ("품사", "pos")
VOCAB_EXTS = (".xlsx", ".xlsm", ".csv")
LETTER_GRADES = {"A": 1, "B": 2, "C": 3, "초급": 1, "중급": 2, "고급": 3}
CACHE_VERSION = 2

# kiwi 품사 → 사전 품사 이름(부분 일치)
KIWI_TO_DICT_POS = {
    "NNG": ("명사",), "NNP": ("명사",), "NNB": ("의존", "명사"), "NR": ("수사", "명사"),
    "NP": ("대명사",), "VV": ("동사",), "VA": ("형용사",), "VX": ("보조", "동사", "형용사"),
    "VCN": ("형용사",), "MAG": ("부사",), "MAJ": ("부사",), "MM": ("관형사",), "IC": ("감탄사",),
    "XR": ("명사", "어근", "부사"),
}


def parse_grade(value):
    """'1등급', '1', 1, 'A등급', '초급' 등을 정수 등급으로 바꾼다."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value) if value == value else None  # NaN 제외
    s = str(value).strip()
    m = re.search(r"\d+", s)
    if m:
        return int(m.group())
    for key, g in LETTER_GRADES.items():
        if s.upper().startswith(key):
            return g
    return None


class Vocabulary:
    def __init__(self, entries: dict[str, list[tuple[int, str]]], source: str = ""):
        self.entries = entries
        self.source = source

    def __len__(self):
        return len(self.entries)

    def __contains__(self, word):
        return word in self.entries

    def grade(self, word: str, kiwi_tag: str | None = None, use_pos: bool = True):
        items = self.entries.get(word)
        if not items:
            return None
        if use_pos and kiwi_tag:
            keys = KIWI_TO_DICT_POS.get(kiwi_tag)
            if keys:
                matched = [g for g, pos in items if any(k in pos for k in keys)]
                if matched:
                    return min(matched)
        return min(g for g, _ in items)


# ----------------------------------------------------------------------
# 파일 찾기
# ----------------------------------------------------------------------
def _find_col(header, names):
    norm = [normalize_header(h).lower() for h in header]
    for name in names:
        for i, h in enumerate(norm):
            if h == name.lower():
                return i
    for name in names:  # 부분 일치 (예: '등급 ', '어휘(표제어)')
        for i, h in enumerate(norm):
            if name.lower() in h:
                return i
    return None


def _iter_tables(path: Path, max_rows=None):
    """(시트 이름, 행 반복자) 목록. 각 행은 값의 튜플."""
    if path.suffix.lower() == ".csv":
        rows = list(csv.reader(read_text(path).splitlines()))
        yield path.stem, iter(rows[:max_rows] if max_rows else rows)
        return
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws.title, ws.iter_rows(values_only=True, max_row=max_rows)
    finally:
        wb.close()


def _locate_header(rows, scan=15):
    """처음 몇 행 안에서 '어휘'와 '등급' 열이 있는 머리글 행을 찾는다."""
    buffered = []
    for _ in range(scan):
        try:
            row = next(rows)
        except StopIteration:
            break
        buffered.append(row)
        wi, gi = _find_col(row, WORD_HEADERS), _find_col(row, GRADE_HEADERS)
        if wi is not None and gi is not None and wi != gi:
            return row, wi, gi, _find_col(row, POS_HEADERS)
    return None


def looks_like_vocab_file(path: Path) -> bool:
    try:
        for _, rows in _iter_tables(path, max_rows=15):
            if _locate_header(rows):
                return True
    except Exception:
        return False
    return False


def find_vocab_file(search_dirs) -> Path | None:
    """앞 폴더부터 어휘 등급 파일을 찾는다. 이름에 '어휘'/'등급'이 들어간 파일을 우선한다."""
    for d in map(Path, search_dirs):
        if not d.is_dir():
            continue
        cands = [p for p in d.iterdir()
                 if p.suffix.lower() in VOCAB_EXTS and not p.name.startswith(("~$", "ERI_", "."))]
        cands.sort(key=lambda p: (not re.search("어휘|등급|vocab", p.name, re.I), p.name))
        for p in cands:
            if looks_like_vocab_file(p):
                return p
    return None


# ----------------------------------------------------------------------
# 불러오기
# ----------------------------------------------------------------------
def _cache_path(path: Path) -> Path:
    return path.with_name(f".{path.name}.eri_cache")


def load_vocabulary(path: str | Path, use_cache: bool = True) -> Vocabulary:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"어휘 등급 파일을 찾을 수 없습니다: {path}")
    stat = path.stat()
    key = (CACHE_VERSION, stat.st_size, int(stat.st_mtime))
    cache = _cache_path(path)
    if use_cache and cache.is_file():
        try:
            with open(cache, "rb") as f:
                saved_key, entries = pickle.load(f)
            if saved_key == key:
                return Vocabulary(entries, str(path))
        except Exception:
            pass

    entries: dict[str, set] = defaultdict(set)
    used_sheets = []
    for sheet, rows in _iter_tables(path):
        found = _locate_header(rows)
        if not found:
            continue
        _, wi, gi, pi = found
        n = 0
        for row in rows:
            if row is None or len(row) <= max(wi, gi):
                continue
            word, grade = row[wi], parse_grade(row[gi])
            if word is None or grade is None:
                continue
            word = re.sub(r"\d+$", "", str(word).strip())  # '가게01' 같은 동형어 번호 제거
            word = word.replace("-", "").replace("^", "").strip()
            if not word:
                continue
            pos = str(row[pi]).strip() if pi is not None and pi < len(row) and row[pi] else ""
            entries[word].add((grade, pos))
            n += 1
        if n:
            used_sheets.append(f"{sheet}({n:,}행)")
    if not entries:
        raise ValueError(f"'{path.name}'에서 '어휘'와 '등급' 열을 찾지 못했습니다.")
    final = {w: sorted(v) for w, v in entries.items()}
    if use_cache:
        try:
            with open(cache, "wb") as f:
                pickle.dump((key, final), f)
        except OSError:
            pass
    vocab = Vocabulary(final, str(path))
    vocab.sheets = used_sheets
    return vocab
